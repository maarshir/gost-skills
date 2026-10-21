"""Флаг --fix: исправленная копия рядом с исходным файлом."""

import json
from pathlib import Path

from docx.oxml.ns import qn
from docx.shared import Mm, Pt

import check_docx
from docbuilder import good_document


def codes(report):
    return {i.code for i in report.issues}


def check(path, config=None):
    return check_docx.Checker(Path(path), check_docx.load_config(config)).run()


def broken():
    b = good_document()
    section = b.doc.sections[0]
    section.left_margin, section.right_margin = Mm(20), Mm(10)
    section.top_margin = Mm(25)
    normal = b.doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    return b


def test_fix_writes_copy_that_passes(tmp_path, capsys):
    path = Path(broken().save(tmp_path / "work.docx"))
    before = path.read_bytes()
    assert check_docx.main([str(path), "--fix"]) == 0
    out = capsys.readouterr().out
    assert path.read_bytes() == before
    copy = tmp_path / "work.fixed.docx"
    assert copy.exists()
    assert "ИСПРАВЛЕНО" in out and "ОСТАЛОСЬ ИСПРАВИТЬ ВРУЧНУЮ" in out
    assert "Поле левое 20 мм → 30 мм" in out
    assert "Шрифт стиля «Обычный»: Calibri → Times New Roman" in out
    assert "Размер шрифта стиля «Обычный»: 11 пт → 12 пт" in out
    report = check(copy)
    assert not codes(report) & {"margin", "font", "size-small"}, [(i.code, i.message) for i in report.issues]


def test_fix_uses_profile_size(tmp_path):
    cfg = tmp_path / "m.toml"
    cfg.write_text("[text]\nsize_pt = 14\n", encoding="utf-8")
    b = good_document()
    b.doc.styles["Normal"].font.size = Pt(12)
    path = Path(b.save(tmp_path / "work.docx"))
    assert check_docx.main([str(path), "--fix", "--config", str(cfg)]) == 0
    assert "size" not in codes(check(tmp_path / "work.fixed.docx", cfg))


def test_fix_replaces_theme_font(tmp_path):
    b = good_document()
    fonts = b.doc.styles["Normal"].element.rPr.rFonts
    fonts.set(qn("w:asciiTheme"), "minorHAnsi")
    fonts.set(qn("w:hAnsiTheme"), "minorHAnsi")
    path = Path(b.save(tmp_path / "work.docx"))
    assert "font" in codes(check(path))
    check_docx.main([str(path), "--fix"])
    assert "font" not in codes(check(tmp_path / "work.fixed.docx"))


def test_fix_leaves_direct_formatting_for_manual(tmp_path, capsys):
    b = good_document()
    p = [x for x in b.doc.paragraphs if len(x.text) > 60][-1]
    for run in p.runs:
        run.font.name = "Arial"
    path = Path(b.save(tmp_path / "work.docx"))
    assert check_docx.main([str(path), "--fix"]) == 1
    out = capsys.readouterr().out
    assert "ничего: однозначных ошибок для исправления нет" in out
    assert "Шрифт Arial вместо Times New Roman" in out


def test_fix_does_not_overwrite_without_force(tmp_path, capsys):
    path = Path(broken().save(tmp_path / "work.docx"))
    copy = tmp_path / "work.fixed.docx"
    copy.write_bytes(b"old")
    assert check_docx.main([str(path), "--fix"]) == 2
    assert copy.read_bytes() == b"old"
    assert "--force" in capsys.readouterr().out
    assert check_docx.main([str(path), "--fix", "--force"]) == 0
    assert copy.read_bytes() != b"old"


def test_fix_json_lists_fixed(tmp_path, capsys):
    path = Path(broken().save(tmp_path / "work.docx"))
    check_docx.main([str(path), "--fix", "--json"])
    data = json.loads(capsys.readouterr().out)
    assert data["file"] == "work.fixed.docx"
    assert any("Поле правое" in item for item in data["fixed"])
    assert data["errors"] == 0


def test_without_fix_no_copy(tmp_path):
    path = Path(broken().save(tmp_path / "work.docx"))
    assert check_docx.main([str(path)]) == 1
    assert not (tmp_path / "work.fixed.docx").exists()


def replace_text(b, old, new):
    for p in b.doc.paragraphs:
        if p.text == old:
            p.runs[0].text = new
            return p
    raise AssertionError(old)


def split_runs(b, old, parts):
    for p in b.doc.paragraphs:
        if p.text == old:
            run = p.runs[0]
            run.text = parts[0]
            for part in parts[1:]:
                p.add_run(part).bold = run.bold
            return p
    raise AssertionError(old)


def fix(tmp_path, b, *args):
    path = Path(b.save(tmp_path / "work.docx"))
    before = path.read_bytes()
    check_docx.main([str(path), "--fix", *args])
    assert path.read_bytes() == before
    copy = tmp_path / "work.fixed.docx"
    return copy, [p.text for p in check_docx.docx.Document(str(copy)).paragraphs]


def test_fix_figure_dash(tmp_path, capsys):
    b = good_document()
    replace_text(b, "Рисунок 1 – Схема поиска", "Рисунок 1 - Схема поиска")
    copy, texts = fix(tmp_path, b)
    assert "Рисунок 1 – Схема поиска" in texts
    assert "Подпись «Рисунок 1 - Схема поиска» → «Рисунок 1 – Схема поиска»" in capsys.readouterr().out
    assert "figure-dash" not in codes(check(copy))


def test_fix_table_dash_after_dot(tmp_path):
    b = good_document()
    replace_text(b, "Таблица 1 – Результаты сравнения", "Таблица 1. Результаты сравнения")
    copy, texts = fix(tmp_path, b)
    assert "Таблица 1 – Результаты сравнения" in texts
    assert "table-dash" not in codes(check(copy))


def test_fix_caption_dot(tmp_path):
    b = good_document()
    replace_text(b, "Рисунок 1 – Схема поиска", "Рисунок 1 – Схема поиска.")
    replace_text(b, "Таблица 1 – Результаты сравнения", "Таблица 1 – Результаты сравнения.")
    copy, texts = fix(tmp_path, b)
    assert "Рисунок 1 – Схема поиска" in texts and "Таблица 1 – Результаты сравнения" in texts
    assert not codes(check(copy)) & {"figure-dot", "table-dot"}


def test_fix_dash_and_dot_together_keeps_formatting(tmp_path):
    b = good_document()
    p = replace_text(b, "Рисунок 1 – Схема поиска", "Рисунок 1: Схема поиска.")
    p.runs[0].italic = True
    copy, _ = fix(tmp_path, b)
    doc = check_docx.docx.Document(str(copy))
    fixed = [x for x in doc.paragraphs if x.text == "Рисунок 1 – Схема поиска"]
    assert fixed and fixed[0].runs[0].italic


def test_fix_heading_number_dot(tmp_path, capsys):
    b = good_document()
    replace_text(b, "1.1 Сравнение", "1.1. Сравнение")
    copy, texts = fix(tmp_path, b)
    assert "1.1 Сравнение" in texts
    assert "Заголовок «1.1. Сравнение» → «1.1 Сравнение»" in capsys.readouterr().out
    assert "heading-number-dot" not in codes(check(copy))


def test_fix_skips_split_runs(tmp_path, capsys):
    b = good_document()
    split_runs(b, "Рисунок 1 – Схема поиска", ["Рисунок 1 ", "- Схема поиска"])
    split_runs(b, "1.1 Сравнение", ["1.1.", " Сравнение"])
    copy, texts = fix(tmp_path, b)
    assert "Рисунок 1 - Схема поиска" in texts and "1.1. Сравнение" in texts
    out = capsys.readouterr().out
    assert "Подпись «" not in out and "Заголовок «" not in out
    assert {"figure-dash", "heading-number-dot"} <= codes(check(copy))


def test_fix_respects_profile_dashes(tmp_path):
    cfg = tmp_path / "m.toml"
    cfg.write_text('[captions]\ndashes = ["—"]\n', encoding="utf-8")
    b = good_document()
    replace_text(b, "Рисунок 1 – Схема поиска", "Рисунок 1 - Схема поиска")
    replace_text(b, "Таблица 1 – Результаты сравнения", "Таблица 1 — Результаты сравнения")
    copy, texts = fix(tmp_path, b, "--config", str(cfg))
    assert "Рисунок 1 — Схема поиска" in texts
    assert not codes(check(copy, cfg)) & {"figure-dash", "table-dash"}
