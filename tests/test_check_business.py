"""Проверка деловых документов: правильное письмо собирается кодом, тесты портят в нём одно место."""

import importlib.util
import json
import re
import sys
from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.shared import Cm, Mm, Pt

SKILL = Path(__file__).resolve().parent.parent / "skills" / "gost-business"
_spec = importlib.util.spec_from_file_location("gost_business_check_docx", SKILL / "scripts" / "check_docx.py")
cb = importlib.util.module_from_spec(_spec)
sys.modules[_spec.name] = cb
_spec.loader.exec_module(cb)

BODY = ("Просим рассмотреть возможность поставки бумаги для офисной техники в объёме двухсот пачек "
        "до 01.04.2026 на условиях договора от 10.01.2026 № 5.")


class Letter:
    """Письмо на угловом бланке по образцу из references/documents.md."""

    def __init__(self) -> None:
        self.doc = Document()
        normal = self.doc.styles["Normal"]
        normal.font.name = "Times New Roman"
        normal.font.size = Pt(14)
        pf = normal.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
        pf.space_after = Pt(0)
        s = self.doc.sections[0]
        s.page_width, s.page_height = Mm(210), Mm(297)
        s.left_margin, s.right_margin = Mm(20), Mm(10)
        s.top_margin, s.bottom_margin = Mm(20), Mm(20)
        for line in ("ООО «Пример»", "ул. Ленина, д. 1, г. Москва, 101000", "05.03.2026 № 12-04/37"):
            self.doc.add_paragraph(line)
        self.doc.add_paragraph("Директору ООО «Партнёр»").alignment = WD_ALIGN_PARAGRAPH.RIGHT
        self.doc.add_paragraph("И.И. Иванову").alignment = WD_ALIGN_PARAGRAPH.RIGHT
        self.doc.add_paragraph("О поставке бумаги")
        self.doc.add_paragraph("Уважаемый Иван Иванович!").alignment = WD_ALIGN_PARAGRAPH.CENTER
        body = self.doc.add_paragraph(BODY)
        body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        body.paragraph_format.first_line_indent = Cm(1.25)
        self.doc.add_paragraph("Приложение: прайс-лист на 2 л. в 1 экз.")
        self.doc.add_paragraph("Генеральный директор\tП.П. Петров")
        self.doc.add_paragraph("")
        for line in ("Сидорова Анна Петровна", "+7 495 123-45-67"):
            self.doc.add_paragraph(line).runs[0].font.size = Pt(10)

    def para(self, start: str):
        return next(p for p in self.doc.paragraphs if p.text.startswith(start))

    def save(self, path) -> str:
        self.doc.save(str(path))
        return str(path)


def check(letter, tmp_path, config=None):
    path = letter.save(tmp_path / "letter.docx")
    return cb.Checker(Path(path), cb.load_config(config)).run()


def codes(report, severity=None):
    return {i.code for i in report.issues if severity is None or i.severity == severity}


def with_text(start, new):
    def spoil(letter):
        p = letter.para(start)
        p.runs[0].text = new
    return spoil


def font(name):
    def spoil(letter):
        letter.para("Просим").runs[0].font.name = name
    return spoil


def size(pt, start="Просим"):
    def spoil(letter):
        letter.para(start).runs[0].font.size = Pt(pt)
    return spoil


def indent(cm):
    def spoil(letter):
        letter.para("Просим").paragraph_format.first_line_indent = Cm(cm)
    return spoil


def margin(name, mm):
    def spoil(letter):
        setattr(letter.doc.sections[0], name, Mm(mm))
    return spoil


def page(w, h):
    def spoil(letter):
        s = letter.doc.sections[0]
        s.page_width, s.page_height = Mm(w), Mm(h)
    return spoil


# Один испорченный элемент на каждый код
SPOILED = {
    "page-size": (cb.ERROR, page(216, 279)),
    "margin": (cb.ERROR, margin("left_margin", 15)),
    "font": (cb.ERROR, font("Courier New")),
    "size-small": (cb.ERROR, size(11)),
    "size-large": (cb.WARNING, size(16)),
    "indent": (cb.ERROR, indent(0)),
    "date-format": (cb.ERROR, with_text("05.03.2026", "5.3.2026 № 12-04/37")),
    "date-invalid": (cb.ERROR, with_text("05.03.2026", "31.02.2026 № 12-04/37")),
    "date-year": (cb.WARNING, with_text("05.03.2026", "5 марта 2026 года № 12-04/37")),
    "addressee-case": (cb.ERROR, with_text("Директору", "Директор ООО «Партнёр»")),
    "title-dot": (cb.ERROR, with_text("О поставке", "О поставке бумаги.")),
    "signature-initials": (cb.ERROR, with_text("Генеральный", "Генеральный директор\tПетров П.П.")),
    "attachment-sheets": (cb.ERROR, with_text("Приложение", "Приложение: прайс-лист в 1 экз.")),
    "attachment-copies": (cb.WARNING, with_text("Приложение", "Приложение: прайс-лист на 2 л.")),
    "executor-name": (cb.ERROR, with_text("Сидорова", "Сидорова")),
    "executor-phone": (cb.ERROR, with_text("+7", "sidorova@primer.ru")),
}


def test_good_letter_has_no_issues(tmp_path):
    report = check(Letter(), tmp_path)
    assert report.issues == [], [(i.code, i.message) for i in report.issues]


@pytest.mark.parametrize("code", sorted(SPOILED))
def test_one_spoiled_place_is_found(code, tmp_path):
    severity, spoil = SPOILED[code]
    letter = Letter()
    spoil(letter)
    report = check(letter, tmp_path)
    assert codes(report) == {code}, [(i.code, i.message) for i in report.issues]
    assert codes(report, severity) == {code}


def test_date_words_only_when_forbidden(tmp_path):
    letter = Letter()
    with_text("05.03.2026", "5 марта 2026 г. № 12-04/37")(letter)
    assert codes(check(letter, tmp_path)) == set()
    cfg = tmp_path / "cfg.toml"
    cfg.write_text("[date]\nwords_allowed = false\n", encoding="utf-8")
    report = check(letter, tmp_path, cfg)
    assert codes(report) == {"date-words"}
    assert "05.03.2026" in report.issues[0].message


@pytest.mark.parametrize("text, good", [
    ("05-03-2026", "05.03.2026"),
    ("05/03/2026", "05.03.2026"),
    ("05.03.26", "05.03.2026"),
    ("2026-03-05", "05.03.2026"),
])
def test_date_format_suggests_fix(text, good, tmp_path):
    letter = Letter()
    with_text("05.03.2026", f"{text} № 12-04/37")(letter)
    report = check(letter, tmp_path)
    assert codes(report) == {"date-format"}
    assert report.issues[0].message.endswith(f"нужно {good}")


def test_numbers_are_not_dates(tmp_path):
    letter = Letter()
    with_text("Генеральный", "См. пункт 1.2.10 и раздел 3.4 договора")(letter)
    assert codes(check(letter, tmp_path)) == set()


def test_margins_may_be_wider_and_a5_is_fine(tmp_path):
    letter = Letter()
    margin("left_margin", 30)(letter)
    page(148, 210)(letter)
    assert codes(check(letter, tmp_path)) == set()


def test_executor_mark_may_be_small_but_body_may_not(tmp_path):
    letter = Letter()
    for line in ("Сидорова", "+7"):
        letter.para(line).runs[0].font.size = Pt(8)
    assert "size-small" not in codes(check(letter, tmp_path))
    letter.para("+7").runs[0].text = "Сидорова Анна Петровна"
    assert "size-small" in codes(check(letter, tmp_path))


def test_font_from_style_and_place(tmp_path):
    letter = Letter()
    letter.doc.styles["Normal"].font.name = "Courier New"
    report = check(letter, tmp_path)
    issue = report.issues[0]
    assert issue.message == "Шрифт Courier New, нужен один из: Times New Roman, Arial, Calibri, Verdana"
    assert len(issue.places) > 5
    assert issue.places[0].paragraph == 1 and issue.places[0].text == "ООО «Пример»"


def test_short_requisites_need_no_indent(tmp_path):
    letter = Letter()
    letter.para("О поставке").paragraph_format.first_line_indent = Cm(0)
    letter.para("Генеральный").paragraph_format.first_line_indent = Cm(2)
    assert codes(check(letter, tmp_path)) == set()


def test_instruction_example_loads(tmp_path):
    cfg = SKILL / "assets" / "instruction-example.toml"
    report = check(Letter(), tmp_path, cfg)
    assert codes(report) == set()
    letter = Letter()
    font("Arial")(letter)
    assert codes(check(letter, tmp_path, cfg)) == {"font"}


def test_cli_text_json_and_exit_codes(tmp_path, capsys):
    good = Letter().save(tmp_path / "good.docx")
    assert cb.main([good]) == 0
    assert "Нарушений не найдено" in capsys.readouterr().out
    letter = Letter()
    margin("right_margin", 5)(letter)
    bad = letter.save(tmp_path / "bad.docx")
    assert cb.main([bad]) == 1
    assert "Поле правое 5 мм, нужно не меньше 10 мм (раздел 4)" in capsys.readouterr().out
    assert cb.main([bad, "--json"]) == 1
    data = json.loads(capsys.readouterr().out)
    assert data["profile"] == "ГОСТ Р 7.0.97-2016" and data["issues"][0]["code"] == "margin"


def test_bad_inputs(tmp_path, capsys):
    assert cb.main([str(tmp_path / "letter.doc")]) == 2
    assert cb.main([str(tmp_path / "nope.docx")]) == 2
    broken = tmp_path / "broken.docx"
    broken.write_text("не архив")
    assert cb.main([str(broken)]) == 2


def insert_before(letter, start, text, right=False):
    p = letter.para(start).insert_paragraph_before(text)
    if right:
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    return p


@pytest.mark.parametrize("word, good", [
    ("Директор", "Директору"), ("Генеральный", "Генеральному"), ("Заместитель", "Заместителю"),
    ("Главный", "Главному"), ("Управляющий", "Управляющему"), ("Технический", "Техническому"),
    ("Глава", "Главе"), ("Начальник", "Начальнику"),
])
def test_dative_suggestion(word, good):
    assert cb.dative(word) == good


def test_addressee_message_and_place(tmp_path):
    letter = Letter()
    with_text("Директору", "Генеральный директор ООО «Партнёр»")(letter)
    issue = check(letter, tmp_path).issues[0]
    assert issue.message == "Адресат в именительном падеже: «Генеральный», нужно «Генеральному»"
    assert issue.places[0].paragraph == 4


def test_addressee_by_indent_or_tab(tmp_path):
    letter = Letter()
    p = letter.para("Директору")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent = Cm(9)
    p.runs[0].text = "Начальник отдела кадров"
    assert codes(check(letter, tmp_path)) == {"addressee-case"}
    p.paragraph_format.left_indent = Cm(0)
    p.runs[0].text = "\tНачальник отдела кадров"
    assert codes(check(letter, tmp_path)) == {"addressee-case"}


def test_doubtful_head_lines_are_skipped(tmp_path):
    letter = Letter()
    # должность автора на бланке слева, гриф утверждения справа, организация адресата
    insert_before(letter, "ООО «Пример»", "Генеральный директор ООО «Пример»")
    for line in ("УТВЕРЖДАЮ", "Генеральный директор", "П.П. Петров"):
        insert_before(letter, "Директору", line, right=True)
    insert_before(letter, "И.И. Иванову", "ООО «Партнёр»", right=True)
    assert codes(check(letter, tmp_path)) == set()


def test_title_may_end_with_abbreviation(tmp_path):
    letter = Letter()
    with_text("О поставке", "Об итогах работы за 2025 г.")(letter)
    assert codes(check(letter, tmp_path)) == set()


def test_signature_on_two_lines(tmp_path):
    letter = Letter()
    with_text("Генеральный", "Генеральный директор")(letter)
    p = letter.para("Генеральный")
    p._p.addnext(letter.doc.add_paragraph("Петров П.П.")._p)
    report = check(letter, tmp_path)
    assert codes(report) == {"signature-initials"}
    assert report.issues[0].message == "Расшифровка подписи «Петров П.П.», нужно «П.П. Петров»"


def test_names_in_text_and_executor_are_not_signatures(tmp_path):
    letter = Letter()
    insert_before(letter, "Генеральный", "Ответственный исполнитель Иванов И.И.")
    with_text("Сидорова", "Сидорова А.П.")(letter)
    assert codes(check(letter, tmp_path)) == set()


def test_signature_check_can_be_turned_off(tmp_path):
    letter = Letter()
    SPOILED["signature-initials"][1](letter)
    cfg = tmp_path / "cfg.toml"
    cfg.write_text("[signature]\ninitials_first = false\n", encoding="utf-8")
    assert codes(check(letter, tmp_path, cfg)) == set()


def test_every_code_is_in_checks_md():
    source = (SKILL / "scripts" / "check_docx.py").read_text(encoding="utf-8")
    in_script = set(re.findall(r"(?:ERROR|WARNING),\s*\"([a-z\-]+)\"", source))
    checks = (SKILL / "references" / "checks.md").read_text(encoding="utf-8")
    documented = set(re.findall(r"^\| `([a-z\-]+)` \|", checks, re.M))
    assert in_script == documented
    assert set(SPOILED) | {"date-words"} == in_script, "у каждого кода свой тест"


def test_numbered_attachments_are_checked_one_by_one(tmp_path):
    letter = Letter()
    with_text("Приложение", "Приложения:")(letter)
    insert_before(letter, "Генеральный", "1. Прайс-лист на 2 л. в 1 экз.")
    insert_before(letter, "Генеральный", "2. Договор поставки в 2 экз.")
    report = check(letter, tmp_path)
    assert codes(report) == {"attachment-sheets"}
    assert report.issues[0].places[0].text.startswith("2. Договор")
    assert report.issues[0].message == "В отметке о приложении нет числа листов"


@pytest.mark.parametrize("text", [
    "Приложение: каталог продукции, сброшюрованный, в 1 экз.",
    "Приложение: письмо ООО «Партнёр» от 01.03.2026 № 15 и приложение к нему, всего на 5 л.",
    "Приложение: 1. Прайс-лист на 2 л. в 1 экз.",
])
def test_doubtful_attachments_are_skipped(text, tmp_path):
    letter = Letter()
    with_text("Приложение", text)(letter)
    assert codes(check(letter, tmp_path)) == set()


def test_attachment_in_text_is_not_a_mark(tmp_path):
    letter = Letter()
    with_text("Приложение", "Приложение к письму направим отдельно")(letter)
    assert codes(check(letter, tmp_path)) == set()


def test_executor_forms_and_missing_mark(tmp_path):
    for name in ("А.П. Сидорова", "Сидорова А.П.", "Исп. Сидорова Анна Петровна"):
        letter = Letter()
        with_text("Сидорова", name)(letter)
        assert codes(check(letter, tmp_path)) == set(), name
    letter = Letter()
    for p in letter.doc.paragraphs[-3:]:
        p._p.getparent().remove(p._p)
    assert codes(check(letter, tmp_path)) == set()


def test_attachment_and_executor_checks_can_be_turned_off(tmp_path):
    letter = Letter()
    SPOILED["attachment-sheets"][1](letter)
    SPOILED["executor-name"][1](letter)
    cfg = tmp_path / "cfg.toml"
    cfg.write_text("[attachment]\ncheck = false\n[executor]\ncheck = false\n", encoding="utf-8")
    assert codes(check(letter, tmp_path, cfg)) == set()
