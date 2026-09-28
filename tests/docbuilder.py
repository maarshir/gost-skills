"""Сборка тестовых .docx прямо в тестах: чужие работы в репозиторий не кладём.

good_document() собирает курсовую, оформленную по ГОСТ 7.32. Тесты портят
в ней одно место и проверяют, что скрипт находит именно его.
"""

from __future__ import annotations

import struct
import zlib
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt


def tiny_png() -> BytesIO:
    """Картинка 2×2 без Pillow: заголовок PNG, IHDR, IDAT и IEND вручную."""
    def chunk(kind: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data))
    raw = b"".join(b"\x00" + b"\x80\x80\x80" * 2 for _ in range(2))
    png = (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 2, 2, 8, 2, 0, 0, 0))
           + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))
    return BytesIO(png)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, " PAGE "), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


class Builder:
    def __init__(self) -> None:
        self.doc = Document()
        normal = self.doc.styles["Normal"]
        normal.font.name = "Times New Roman"
        normal.font.size = Pt(14)
        pf = normal.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        pf.first_line_indent = Cm(1.25)
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.space_after = Pt(0)
        section = self.doc.sections[0]
        section.page_width, section.page_height = Mm(210), Mm(297)
        section.left_margin, section.right_margin = Mm(30), Mm(15)
        section.top_margin, section.bottom_margin = Mm(20), Mm(20)
        section.different_first_page_header_footer = True
        footer_p = section.footer.paragraphs[0]
        footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_page_field(footer_p)

    # --- кирпичики

    def centered(self, text: str, bold: bool = False, new_page: bool = False):
        p = self.doc.add_paragraph()
        p.add_run(text).bold = bold
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.page_break_before = new_page
        return p

    def structural(self, text: str, **kw):
        return self.centered(text, bold=True, new_page=kw.get("new_page", True))

    def heading(self, text: str, new_page: bool = False):
        p = self.doc.add_paragraph()
        p.add_run(text).bold = True
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.page_break_before = new_page
        return p

    def text(self, text: str):
        return self.doc.add_paragraph(text)

    def picture(self):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.add_run().add_picture(tiny_png(), width=Cm(3))
        return p

    def figure_caption(self, text: str):
        return self.centered(text)

    def table_title(self, text: str):
        p = self.doc.add_paragraph(text)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.first_line_indent = Cm(0)
        return p

    def table(self):
        t = self.doc.add_table(rows=2, cols=2)
        t.cell(0, 0).text, t.cell(0, 1).text = "Метод", "Точность"
        t.cell(1, 0).text, t.cell(1, 1).text = "А", "0,9"
        return t

    def save(self, path) -> str:
        self.doc.save(str(path))
        return str(path)


def good_document() -> Builder:
    b = Builder()
    b.centered("Министерство науки и высшего образования")
    b.centered("КУРСОВАЯ РАБОТА", bold=True)
    b.centered("Тема: Поиск по документам")
    b.structural("СОДЕРЖАНИЕ")
    for line in ("Введение\t3", "1 Обзор методов\t4", "Заключение\t6"):
        b.text(line)
    b.structural("ВВЕДЕНИЕ")
    b.text("Работа посвящена поиску ответов в документах [1]. Цель работы состоит в сравнении методов.")
    b.heading("1 Обзор методов", new_page=True)
    b.text("Схема поиска показана на рисунке 1, подробности в работе [2].")
    b.picture()
    b.figure_caption("Рисунок 1 – Схема поиска")
    b.heading("1.1 Сравнение")
    b.text("Результаты сравнения приведены в таблице 1.")
    b.table_title("Таблица 1 – Результаты сравнения")
    b.table()
    b.text("Метод А оказался точнее на выбранном наборе вопросов.")
    b.structural("ЗАКЛЮЧЕНИЕ")
    b.text("Сравнение показало, что метод А подходит лучше.")
    b.structural("СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ")
    b.text("1. Иванов И. И. Поиск в тексте. Москва : Наука, 2020. 200 с.")
    b.text("2. Петров П. П. Ранжирование документов // Вестник. 2021. № 3. С. 10–20.")
    b.structural("ПРИЛОЖЕНИЕ А")
    b.text("Исходные данные.")
    return b
