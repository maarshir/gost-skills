"""Проверка списка литературы: правильная запись проходит, в испорченной находится одно нарушение.

Каждый тест берёт исправную запись, портит в ней одно место и смотрит, что скрипт
назвал именно его. Дата «сегодня» зафиксирована, чтобы тесты не старели.
"""

import datetime as dt
import json
import re
from pathlib import Path

import docx
import pytest

import check_bibliography as cb

ROOT = Path(__file__).resolve().parent.parent
SKILL = ROOT / "skills" / "gost-bibliography"
TODAY = dt.date(2026, 9, 30)

BOOK = ("Петров, П. П. Основы информационного поиска : учебное пособие / П. П. Петров, С. С. Сидоров. "
        "– Москва : Наука, 2021. – 256 с. – ISBN 978-5-02-040000-1. – Текст : непосредственный.")
ARTICLE = ("Сидоров, С. С. Сравнение методов ранжирования / С. С. Сидоров, А. А. Козлов. "
           "– Текст : непосредственный // Вестник информатики. – 2022. – № 3. – С. 15–20.")
WEB = ("Орлов, О. О. Поиск по документам / О. О. Орлов. – Текст : электронный // Блог разработчика : [сайт]. "
       "– 2024. – URL: https://example.com/blog/search-2024 (дата обращения: 15.09.2026).")
LAW = ("Об информации, информационных технологиях и о защите информации : Федеральный закон № 149-ФЗ : "
       "принят Государственной Думой 8 июля 2006 года. – Текст : электронный // Официальный интернет-портал "
       "правовой информации. – URL: http://pravo.gov.ru (дата обращения: 15.09.2026).")


def run(lines, config=None, today=TODAY):
    if isinstance(lines, str):
        lines = [lines]
    return cb.check_entries(lines, config or cb.load_config(), today)


def config_with(**sections):
    return cb.deep_merge(cb.load_config(), sections)


def codes(report, severity=None):
    return {i.code for i in report.issues if severity is None or i.severity == severity}


def examples():
    text = (SKILL / "references" / "examples.md").read_text(encoding="utf-8")
    return [line for block in re.findall(r"```\n(.*?)```", text, re.S) for line in block.splitlines() if line.strip()]


# ------------------------------------------------------------ исправные записи

@pytest.mark.parametrize("entry", [BOOK, ARTICLE, WEB, LAW])
def test_good_entries_are_clean(entry):
    assert codes(run(entry)) == set()


def test_every_example_in_references_has_no_errors():
    """Образцы из examples.md проходят без ошибок: справка и скрипт не противоречат друг другу."""
    lines = examples()
    assert len(lines) >= 12
    report = run(lines)
    assert report.count(cb.ERROR) == 0, cb.render_text("examples", cb.load_config(), report)


# ------------------------------------------------------------ по одному нарушению

BROKEN = [
    ("hyphen-dash", BOOK.replace("2021. – 256", "2021. - 256")),
    ("dash-without-dot", BOOK.replace("2021. – 256", "2021 – 256")),
    ("end-dot", BOOK.rstrip(".")),
    ("double-punct", BOOK.replace("256 с.", "256 с..")),
    ("space-colon", BOOK.replace("Москва : Наука", "Москва: Наука")),
    ("space-slash", BOOK.replace("пособие / П. П.", "пособие/ П. П.")),
    ("space-slash", ARTICLE.replace("непосредственный // Вестник", "непосредственный//Вестник")),
    ("space-before", BOOK.replace("Наука, 2021", "Наука , 2021")),
    ("heading-comma", BOOK.replace("Петров, П. П. Основы", "Петров П. П. Основы")),
    ("initials-space", BOOK.replace("/ П. П. Петров", "/ П.П. Петров")),
    ("et-al", "Методы : монография / А. А. Иванов, Б. Б. Смирнов, В. В. Кузнецов и др. – Москва : Наука, 2020. – 312 с."),
    ("year-missing", BOOK.replace(", 2021.", ".")),
    ("year-future", BOOK.replace("2021", "2031")),
    ("pages-abbr", BOOK.replace("256 с.", "256 стр.")),
    ("pages-hyphen", ARTICLE.replace("С. 15–20", "С. 15-20")),
    ("pages-space", ARTICLE.replace("С. 15–20", "С.15–20")),
    ("pages-case", ARTICLE.replace("С. 15–20", "с. 15–20")),
    ("pages-case", BOOK.replace("256 с.", "256 С.")),
    ("number-sign", ARTICLE.replace("№ 3", "№3")),
    ("url-prefix", WEB.replace("URL: https", "https")),
    ("url-access-date", WEB.replace(" (дата обращения: 15.09.2026)", "")),
    ("url-date-format", WEB.replace("15.09.2026", "2026-09-15")),
    ("url-date-format", WEB.replace("15.09.2026", "31.02.2026")),
    ("url-date-future", WEB.replace("15.09.2026", "15.10.2026")),
    ("content-type-format", BOOK.replace("Текст : непосредственный", "Текст: непосредственный")),
    ("content-type-format", BOOK.replace("Текст : непосредственный", "Текст : Непосредственный")),
]


@pytest.mark.parametrize("code, entry", BROKEN, ids=[f"{c}-{i}" for i, (c, _) in enumerate(BROKEN)])
def test_one_broken_place_is_found(code, entry):
    report = run(entry)
    assert code in codes(report, cb.ERROR), cb.render_text("t", cb.load_config(), report)


@pytest.mark.parametrize("code, entry", [
    ("no-area-dash", "Петров, П. П. Основы поиска / П. П. Петров. Москва : Наука, 2021. 256 с."),
    ("authors-many", "Методы : монография / А. А. Иванов, Б. Б. Смирнов, В. В. Кузнецов, Г. Г. Орлов, "
                     "Д. Д. Попов. – Москва : Наука, 2020. – 312 с."),
    ("heading-four", "Иванов, А. А. Анализ текстов / А. А. Иванов, Б. Б. Смирнов, В. В. Кузнецов, Г. Г. Орлов. "
                     "– Москва : Наука, 2022. – 320 с."),
    ("volume-missing", BOOK.replace(" – 256 с.", "")),
    ("pages-missing", ARTICLE.replace(" – С. 15–20.", "")),
    ("old-electronic", WEB.replace("Поиск по документам", "Поиск по документам [Электронный ресурс]")),
    ("old-access", WEB.replace("URL:", "Режим доступа:")),
    ("content-type-mismatch", WEB.replace("электронный", "непосредственный")),
], ids=lambda v: v if len(v) < 30 else "")
def test_warnings(code, entry):
    report = run(entry)
    assert code in codes(report, cb.WARNING)
    assert report.count(cb.ERROR) == 0 or code == "old-access"


def test_site_without_year_is_only_a_warning():
    entry = "Документация библиотеки. – URL: https://example.com/docs/ (дата обращения: 15.09.2026). – Текст : электронный."
    assert codes(run(entry), cb.WARNING) == {"year-missing"}


def test_dash_in_title_is_a_warning_not_an_error():
    entry = BOOK.replace("информационного поиска", "поиска – обзор")
    report = run(entry)
    assert "dash-without-dot" in codes(report, cb.WARNING)
    assert report.count(cb.ERROR) == 0


def test_url_is_not_checked_for_colons_and_slashes():
    entry = WEB.replace("search-2024", "a/b:c//d")
    assert codes(run(entry)) == set()


def test_isbn_and_access_date_are_not_years():
    report = run(BOOK.replace("ISBN 978-5-02-040000-1", "ISBN 978-5-2035-2040-1"))
    assert "year-future" not in codes(report)
    assert "year-future" not in codes(run(WEB.replace("2024", "2025"), today=dt.date(2026, 1, 1)))


def test_edition_abbreviations_are_not_double_punctuation():
    """«изд., перераб. и доп.» это сокращения, а не лишние знаки."""
    assert codes(run(BOOK.replace("Сидоров. – Москва", "Сидоров. – 2-е изд., перераб. и доп. – Москва"))) == set()


# ------------------------------------------------------------ настройки методички

def test_place_abbreviation_only_with_setting():
    entry = BOOK.replace("Москва :", "М. :")
    assert "place-abbr" not in codes(run(entry))
    assert "place-abbr" in codes(run(entry, config_with(place={"allow_abbreviations": False})), cb.ERROR)


def test_content_type_required():
    entry = BOOK.replace(" – Текст : непосредственный.", "")
    assert codes(run(entry)) == set()
    cfg = config_with(content_type={"required": True})
    assert "content-type-missing" in codes(run(entry, cfg), cb.ERROR)
    # законы и стандарты не требуют
    assert "content-type-missing" not in codes(run(LAW.replace(" – Текст : электронный", ""), cfg))


def test_separator_dot_allowed():
    entry = "Петров, П. П. Основы поиска / П. П. Петров. Москва : Наука, 2021. 256 с."
    assert "no-area-dash" not in codes(run(entry, config_with(areas={"separator": "dot"})))
    assert "hyphen-dash" not in codes(run(BOOK.replace(". – 256", ". - 256"), config_with(areas={"separator": "any"})))


def test_freshness():
    cfg = config_with(freshness={"max_age_years": 5})
    old = BOOK.replace("2021", "2015")
    assert "stale" in codes(run(old, cfg), cb.WARNING)
    assert "stale" not in codes(run(BOOK, cfg))
    assert "stale" not in codes(run(LAW, cfg)), "законы не устаревают по этой настройке"


def test_methodichka_example_loads():
    cfg = cb.load_config(SKILL / "assets" / "methodichka-example.toml")
    assert cfg["order"]["mode"] == "alphabet" and cfg["content_type"]["required"] is True
    assert cfg["areas"]["separator"] == "dash", "не заданное в методичке берётся из ГОСТ"


# ------------------------------------------------------------ весь список

def test_numbering():
    report = run([f"1. {BOOK}", f"3. {ARTICLE}"])
    issue = next(i for i in report.issues if i.code == "numbering")
    assert [p.entry for p in issue.places] == [2]
    assert codes(run([f"1. {BOOK}", f"2) {ARTICLE}", f"[3] {WEB}"])) == set()


def test_duplicate_by_text_and_by_url():
    assert "duplicate" in codes(run([BOOK, BOOK.replace(" – ", " — ")]))
    other = WEB.replace("Орлов, О. О. Поиск по документам / О. О. Орлов", "Орлов, О. О. Поиск / О. О. Орлов")
    report = run([WEB, other])
    issue = next(i for i in report.issues if i.code == "duplicate")
    assert issue.places[0].entry == 2 and "запись 1" in issue.places[0].detail


def test_order_alphabet_cyrillic_then_latin():
    latin = "Smith, J. Basics / J. Smith. – London : Example Press, 2019. – 300 p."
    cfg = config_with(order={"mode": "alphabet"})
    assert "order" not in codes(run([ARTICLE, BOOK], config_with(order={"mode": "any"})))
    assert "order" in codes(run([ARTICLE, BOOK], cfg))
    assert "order" not in codes(run([BOOK, ARTICLE, latin], cfg))
    assert "order" in codes(run([latin, BOOK], cfg))


def test_order_legal_first():
    cfg = config_with(order={"mode": "alphabet", "legal_first": True})
    assert "order" not in codes(run([LAW, BOOK, ARTICLE], cfg))
    report = run([BOOK, LAW, ARTICLE], cfg)
    issue = next(i for i in report.issues if i.code == "order")
    assert issue.places[0].entry == 2 and "закон" in issue.places[0].detail


# ------------------------------------------------------------ файлы и вывод

def make_docx(path, entries, numbered=False):
    d = docx.Document()
    d.add_paragraph("СОДЕРЖАНИЕ")
    d.add_paragraph("Список использованных источников")  # строка содержания, не раздел
    d.add_paragraph("ВВЕДЕНИЕ")
    d.add_paragraph("Текст работы со ссылкой [1].")
    d.add_paragraph("СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ")
    for e in entries:
        d.add_paragraph(e, style="List Number" if numbered else None)
    d.add_paragraph("")
    d.add_paragraph("ПРИЛОЖЕНИЕ А")
    d.add_paragraph("Текст приложения без точки")
    d.save(path)
    return path


def test_docx_takes_only_the_bibliography(tmp_path):
    path = make_docx(tmp_path / "work.docx", [BOOK, ARTICLE.replace("№ 3", "№3")])
    lines, notes = cb.read_docx(path)
    assert lines == [BOOK, ARTICLE.replace("№ 3", "№3")]
    assert notes == []
    assert cb.main([str(path), "--today", "2026-09-30"]) == 1


def test_docx_auto_numbering_is_noted(tmp_path):
    path = make_docx(tmp_path / "work.docx", [BOOK, ARTICLE], numbered=True)
    _, notes = cb.read_docx(path)
    assert notes and "автоматическая" in notes[0]


def test_docx_without_bibliography(tmp_path, capsys):
    d = docx.Document()
    d.add_paragraph("ВВЕДЕНИЕ")
    d.save(tmp_path / "no.docx")
    assert cb.main([str(tmp_path / "no.docx")]) == 2
    assert "не найден заголовок" in capsys.readouterr().out


def test_text_file_json_and_exit_codes(tmp_path, capsys):
    good = tmp_path / "good.txt"
    good.write_text(f"1. {BOOK}\n\n2. {ARTICLE}\n", encoding="utf-8")
    assert cb.main([str(good), "--today", "2026-09-30"]) == 0
    assert "Нарушений не найдено" in capsys.readouterr().out
    bad = tmp_path / "bad.txt"
    bad.write_text(BOOK.replace("Москва : Наука", "Москва: Наука"), encoding="utf-8")
    assert cb.main([str(bad), "--json", "--today", "2026-09-30"]) == 1
    data = json.loads(capsys.readouterr().out)
    assert data["errors"] == 1 and data["issues"][0]["code"] == "space-colon"
    assert data["issues"][0]["places"][0]["entry"] == 1


def test_bad_inputs(tmp_path, capsys):
    assert cb.main([str(tmp_path / "missing.txt")]) == 2
    empty = tmp_path / "empty.txt"
    empty.write_text("\n\n", encoding="utf-8")
    assert cb.main([str(empty)]) == 2
    assert cb.main([str(tmp_path / "old.doc")]) == 2


def test_readme_example_is_real(capsys):
    """Вывод в README совпадает с настоящим выводом скрипта на examples/bibliography.txt."""
    code = cb.main([str(ROOT / "examples" / "bibliography.txt"), "--today", "2026-09-30"])
    out = capsys.readouterr().out
    assert code == 1
    readme = (ROOT / "README.md").read_text(encoding="utf-8")
    block = readme.split("check_bibliography.py examples/bibliography.txt\n", 1)[1].split("```", 1)[0]
    assert block == out
