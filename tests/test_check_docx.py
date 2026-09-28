import json

import pytest
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.shared import Cm, Mm, Pt, RGBColor

import check_docx
from docbuilder import good_document


def check(builder, tmp_path, config=None):
    path = builder.save(tmp_path / "work.docx")
    cfg = check_docx.load_config(config)
    return check_docx.Checker(check_docx.Path(path), cfg).run()


def codes(report):
    return {i.code for i in report.issues}


def find(report, code):
    return next(i for i in report.issues if i.code == code)


def para(builder, text):
    return next(p for p in builder.doc.paragraphs if p.text == text)


# ------------------------------------------------------------ исправный документ

def test_good_document_has_no_issues(tmp_path):
    report = check(good_document(), tmp_path)
    assert report.issues == [], [(i.code, i.message) for i in report.issues]


def test_cli_exit_codes(tmp_path, capsys):
    good = good_document().save(tmp_path / "good.docx")
    assert check_docx.main([good]) == 0
    b = good_document()
    b.doc.sections[0].left_margin = Mm(20)
    bad = b.save(tmp_path / "bad.docx")
    assert check_docx.main([bad]) == 1
    out = capsys.readouterr().out
    assert "Поле левое 20 мм, нужно 30 мм (п. 6.1.1)" in out


def test_cli_rejects_doc_and_missing_file(tmp_path, capsys):
    assert check_docx.main([str(tmp_path / "work.doc")]) == 2
    assert check_docx.main([str(tmp_path / "nope.docx")]) == 2
    broken = tmp_path / "broken.docx"
    broken.write_text("не архив")
    assert check_docx.main([str(broken)]) == 2


def test_json_output(tmp_path, capsys):
    b = good_document()
    b.doc.sections[0].right_margin = Mm(10)
    path = b.save(tmp_path / "w.docx")
    check_docx.main([path, "--json"])
    data = json.loads(capsys.readouterr().out)
    assert data["errors"] == 1
    assert data["issues"][0]["code"] == "margin"


# ------------------------------------------------------------ страница

def test_margins_and_page_size(tmp_path):
    b = good_document()
    s = b.doc.sections[0]
    s.top_margin, s.page_width, s.page_height = Mm(25), Mm(216), Mm(279)
    report = check(b, tmp_path)
    messages = [i.message for i in report.issues]
    assert "Поле верхнее 25 мм, нужно 20 мм" in messages
    assert any(m.startswith("Размер листа 216×279") for m in messages)


def test_margin_tolerance(tmp_path):
    b = good_document()
    b.doc.sections[0].left_margin = Mm(30.5)  # Word округляет поля до твипов
    assert "margin" not in codes(check(b, tmp_path))


# ------------------------------------------------------------ номера страниц

def test_page_number_missing(tmp_path):
    b = good_document()
    footer = b.doc.sections[0].footer.paragraphs[0]
    footer._p.getparent().remove(footer._p)
    assert "page-number-missing" in codes(check(b, tmp_path))


def test_page_number_not_centered(tmp_path):
    b = good_document()
    b.doc.sections[0].footer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    assert "page-number-align" in codes(check(b, tmp_path))


def test_page_number_on_title_page(tmp_path):
    b = good_document()
    b.doc.sections[0].different_first_page_header_footer = False
    assert "page-number-title" in codes(check(b, tmp_path))


def test_page_number_in_header(tmp_path):
    b = good_document()
    s = b.doc.sections[0]
    footer = s.footer.paragraphs[0]
    footer._p.getparent().remove(footer._p)
    from docbuilder import add_page_field
    add_page_field(s.header.paragraphs[0])
    assert "page-number-top" in codes(check(b, tmp_path))


# ------------------------------------------------------------ основной текст

def test_wrong_font_is_reported_with_place(tmp_path):
    b = good_document()
    text = "Метод А оказался точнее на выбранном наборе вопросов."
    for run in para(b, text).runs:
        run.font.name = "Arial"
    report = check(b, tmp_path)
    issue = find(report, "font")
    assert issue.message == "Шрифт Arial вместо Times New Roman"
    assert issue.places[0].text.startswith("Метод А оказался")
    assert issue.places[0].section == "1.1 Сравнение"


def test_font_from_style_is_resolved(tmp_path):
    b = good_document()
    b.doc.styles["Normal"].font.name = "Calibri"
    issue = find(check(b, tmp_path), "font")
    assert issue.message == "Шрифт Calibri вместо Times New Roman"
    assert len(issue.places) > 3  # весь основной текст, а не одно место


def test_small_font(tmp_path):
    b = good_document()
    for run in para(b, "Исходные данные.").runs:
        run.font.size = Pt(11)
    assert find(check(b, tmp_path), "size-small").message == "Шрифт 11 пт, меньше 12 пт"


def test_colored_text(tmp_path):
    b = good_document()
    for run in para(b, "Исходные данные.").runs:
        run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    assert "color" in codes(check(b, tmp_path))


def test_line_spacing(tmp_path):
    b = good_document()
    para(b, "Исходные данные.").paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    assert find(check(b, tmp_path), "spacing").message == "Междустрочный интервал 1 вместо 1,5"


def test_exact_line_spacing_is_warning(tmp_path):
    b = good_document()
    para(b, "Исходные данные.").paragraph_format.line_spacing = Pt(18)
    issue = find(check(b, tmp_path), "spacing-exact")
    assert issue.severity == check_docx.WARNING


def test_indent(tmp_path):
    b = good_document()
    para(b, "Исходные данные.").paragraph_format.first_line_indent = Cm(1)
    assert find(check(b, tmp_path), "indent").message == "Абзацный отступ 1,00 см вместо 1,25 см"


def test_title_page_is_not_checked_as_text(tmp_path):
    b = good_document()
    for run in para(b, "Министерство науки и высшего образования").runs:
        run.font.name = "Arial"
        run.font.size = Pt(10)
    assert check(b, tmp_path).issues == []


def test_document_without_structural_elements_is_checked(tmp_path):
    from docbuilder import Builder
    b = Builder()
    b.text("Просто текст без заголовков.")
    for run in b.doc.paragraphs[0].runs:
        run.font.name = "Arial"
    report = check(b, tmp_path)
    assert "font" in codes(report)
    assert report.notes and "титульный лист не отделён" in report.notes[0]


def test_empty_lines_instead_of_page_break(tmp_path):
    b = good_document()
    target = para(b, "Исходные данные.")
    for _ in range(4):
        target.insert_paragraph_before("")
    assert "empty-lines" in codes(check(b, tmp_path))


def test_bold_body_paragraph_hints_at_heading(tmp_path):
    b = good_document()
    for run in para(b, "Исходные данные.").runs:
        run.bold = True
    assert "bold-body" in codes(check(b, tmp_path))


# ------------------------------------------------------------ заголовки

def test_structural_title_checks(tmp_path):
    b = good_document()
    p = para(b, "ЗАКЛЮЧЕНИЕ")
    p.runs[0].text = "Заключение."
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.runs[0].bold = False
    p.paragraph_format.page_break_before = False
    found = codes(check(b, tmp_path))
    assert {"title-case", "title-align", "title-dot", "title-bold", "new-page"} <= found


def test_page_break_run_counts_as_new_page(tmp_path):
    b = good_document()
    para(b, "ЗАКЛЮЧЕНИЕ").paragraph_format.page_break_before = False
    para(b, "Метод А оказался точнее на выбранном наборе вопросов.").add_run().add_break(WD_BREAK.PAGE)
    assert "new-page" not in codes(check(b, tmp_path))


def test_caps_property_counts_as_uppercase(tmp_path):
    b = good_document()
    run = para(b, "ЗАКЛЮЧЕНИЕ").runs[0]
    run.text = "Заключение"
    run.font.all_caps = True
    assert "title-case" not in codes(check(b, tmp_path))


def test_missing_required_element(tmp_path):
    b = good_document()
    p = para(b, "ЗАКЛЮЧЕНИЕ")
    p._p.getparent().remove(p._p)
    issue = find(check(b, tmp_path), "missing")
    assert issue.message == "Нет структурного элемента «ЗАКЛЮЧЕНИЕ»"


def test_alias_name_is_warning(tmp_path):
    b = good_document()
    para(b, "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ").runs[0].text = "СПИСОК ЛИТЕРАТУРЫ"
    report = check(b, tmp_path)
    assert "alias" in codes(report)
    assert "missing" not in codes(report)


def test_heading_checks(tmp_path):
    b = good_document()
    p = para(b, "1.1 Сравнение")
    p.runs[0].text = "1.1. сравнение."
    found = codes(check(b, tmp_path))
    assert {"heading-number-dot", "heading-lower", "heading-dot"} <= found


def test_heading_centered(tmp_path):
    b = good_document()
    para(b, "1.1 Сравнение").alignment = WD_ALIGN_PARAGRAPH.CENTER
    assert "heading-center" in codes(check(b, tmp_path))


def test_section_must_start_on_new_page(tmp_path):
    b = good_document()
    para(b, "1 Обзор методов").paragraph_format.page_break_before = False
    assert find(check(b, tmp_path), "new-page").places[0].text == "1 Обзор методов"


def test_heading_sequence(tmp_path):
    b = good_document()
    para(b, "1.1 Сравнение").runs[0].text = "1.3 Сравнение"
    assert find(check(b, tmp_path), "heading-sequence").message == "Номер 1.3 идёт после 1, ожидался 1.1"


def test_config_can_add_own_title(tmp_path):
    b = good_document()
    para(b, "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ").runs[0].text = "СПИСОК ЛИТЕРАТУРЫ"
    cfg = tmp_path / "m.toml"
    cfg.write_text(
        '[headings]\nstructural = ["СОДЕРЖАНИЕ", "ВВЕДЕНИЕ", "ЗАКЛЮЧЕНИЕ", "СПИСОК ЛИТЕРАТУРЫ"]\n'
        'required = ["СОДЕРЖАНИЕ", "ВВЕДЕНИЕ", "ЗАКЛЮЧЕНИЕ", "СПИСОК ЛИТЕРАТУРЫ"]\n',
        encoding="utf-8",
    )
    assert check(b, tmp_path, cfg).issues == []


# ------------------------------------------------------------ рисунки и таблицы

def test_figure_caption_format(tmp_path):
    b = good_document()
    para(b, "Рисунок 1 – Схема поиска").runs[0].text = "Рис. 1 - схема поиска."
    found = codes(check(b, tmp_path))
    assert {"figure-word", "figure-dash", "figure-lower", "figure-dot"} <= found


def test_figure_caption_above_picture(tmp_path):
    b = good_document()
    caption = para(b, "Рисунок 1 – Схема поиска")
    picture = caption._p.getprevious()
    caption._p.addnext(picture)  # теперь подпись над рисунком
    assert "figure-above" in codes(check(b, tmp_path))


def test_reference_in_text_is_not_a_caption(tmp_path):
    b = good_document()
    para(b, "Исходные данные.").runs[0].text = "Рисунок 1 показывает схему, она подробно разобрана выше."
    report = check(b, tmp_path)
    assert "figure-word" not in codes(report)
    assert "numbering" not in codes(report)


def test_figure_without_reference(tmp_path):
    b = good_document()
    para(b, "Схема поиска показана на рисунке 1, подробности в работе [2].").runs[0].text = (
        "Схема поиска описана ниже, подробности в работе [2]."
    )
    assert find(check(b, tmp_path), "figure-noref").message == "В тексте нет ссылки на рисунок 1"


def test_table_title_checks(tmp_path):
    b = good_document()
    p = para(b, "Таблица 1 – Результаты сравнения")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    found = codes(check(b, tmp_path))
    assert "table-align" in found


def test_table_title_with_indent(tmp_path):
    b = good_document()
    para(b, "Таблица 1 – Результаты сравнения").paragraph_format.first_line_indent = Cm(1.25)
    assert "table-indent" in codes(check(b, tmp_path))


def test_table_without_title(tmp_path):
    b = good_document()
    p = para(b, "Таблица 1 – Результаты сравнения")
    p._p.getparent().remove(p._p)
    assert "table-untitled" in codes(check(b, tmp_path))


def test_figure_numbering(tmp_path):
    b = good_document()
    target = para(b, "Метод А оказался точнее на выбранном наборе вопросов.")
    target.runs[0].text = "Итог показан на рисунке 3."
    picture = b.picture()
    caption = b.figure_caption("Рисунок 3 – Итог")
    target._p.addnext(picture._p)
    picture._p.addnext(caption._p)
    issue = find(check(b, tmp_path), "numbering")
    assert issue.message == "Рисунок 3: ожидался номер 2"


# ------------------------------------------------------------ источники и приложения

def test_citation_to_missing_source(tmp_path):
    b = good_document()
    para(b, "Сравнение показало, что метод А подходит лучше.").runs[0].text = "Итог совпадает с [5, с. 12]."
    assert find(check(b, tmp_path), "source-missing").message == "Ссылка [5], а в списке источников их 2"


def test_uncited_source(tmp_path):
    b = good_document()
    p = para(b, "Работа посвящена поиску ответов в документах [1]. Цель работы состоит в сравнении методов.")
    p.runs[0].text = "Работа посвящена поиску ответов в документах [2]."
    q = para(b, "Схема поиска показана на рисунке 1, подробности в работе [2].")
    q.runs[0].text = "Схема поиска показана на рисунке 1."
    report = check(b, tmp_path)
    assert find(report, "source-uncited").message == "На источник 1 нет ссылки в тексте"


def test_sources_order(tmp_path):
    b = good_document()
    para(b, "Работа посвящена поиску ответов в документах [1]. Цель работы состоит в сравнении методов.") \
        .runs[0].text = "Работа посвящена поиску ответов [2], см. также [1]."
    assert "source-order" in codes(check(b, tmp_path))


def test_citation_ranges():
    assert check_docx.Checker._numbers("1, 3–5") == ["1", "3", "4", "5"]
    assert check_docx.Checker._numbers("А.1 и 2") == ["А.1", "2"]


def test_appendix_letters(tmp_path):
    b = good_document()
    b.structural("ПРИЛОЖЕНИЕ З")
    b.text("Ещё данные.")
    issue = find(check(b, tmp_path), "appendix-letter")
    assert issue.severity == check_docx.ERROR


def test_appendix_order(tmp_path):
    b = good_document()
    b.structural("ПРИЛОЖЕНИЕ В")
    b.text("Ещё данные.")
    assert find(check(b, tmp_path), "appendix-order").message == "Приложение В идёт вместо ожидаемого Б"


# ------------------------------------------------------------ настройки методички

def test_methodichka_example_config(tmp_path):
    from pathlib import Path
    example = Path(check_docx.__file__).parent.parent / "assets" / "methodichka-example.toml"
    b = good_document()
    b.doc.styles["Normal"].font.size = Pt(12)
    para(b, "Исходные данные.").alignment = WD_ALIGN_PARAGRAPH.LEFT
    report = check(b, tmp_path, example)
    found = codes(report)
    assert "size" in found and "align" in found
    # без методички 12 пт по ГОСТ допустимы, а выравнивание не проверяется
    assert check(b, tmp_path).issues == []


def test_render_text_limits_places(tmp_path):
    b = good_document()
    b.doc.styles["Normal"].font.name = "Arial"
    report = check(b, tmp_path)
    text = check_docx.render_text(check_docx.Path("w.docx"), check_docx.load_config(), report, limit=2)
    assert "и ещё" in text
    assert text.startswith("Проверка: w.docx, профиль «ГОСТ 7.32-2017»")


@pytest.mark.parametrize("value,expected", [("1", True), ("0", False), ("false", False), (None, True)])
def test_is_on(value, expected):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    el = OxmlElement("w:b")
    if value is not None:
        el.set(qn("w:val"), value)
    assert check_docx.is_on(el) is expected


def test_example_from_readme(tmp_path):
    """Учебный файл из examples/ даёт ровно те ошибки, что показаны в README."""
    import importlib.util
    from pathlib import Path
    spec = importlib.util.spec_from_file_location(
        "make_example", Path(__file__).resolve().parent.parent / "examples" / "make_example.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    path = module.build(str(tmp_path / "example.docx"))
    report = check_docx.Checker(check_docx.Path(path), check_docx.load_config()).run()
    assert codes(report) == {"margin", "title-align", "new-page", "heading-number-dot", "font",
                             "figure-word", "figure-dash", "figure-dot", "source-missing"}
    assert report.count(check_docx.WARNING) == 0
